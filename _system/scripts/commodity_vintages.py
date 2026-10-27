"""Evaluate archived publisher forecasts without fitting to later outcomes."""
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date,datetime
import io
import math
import re
from urllib.parse import urljoin

from commodity_common import digest,finite,load,save

ARCHIVE='https://www.eia.gov/outlooks/steo/outlook.php'


def parse(raw):
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    dates=list(book['Dates'].iter_rows(values_only=True))
    calendar=next(r for r in dates if any('AAAA_DATEX' in str(v) for v in r))
    prices=list(book['2tab'].iter_rows(values_only=True))
    if not any('dollars per barrel' in str(v).lower() for r in prices[:8] for v in r):raise ValueError('Unexpected oil price units')
    result={}
    for sid,name in [('WTIPUUS','wti_spot'),('BREPUUS','brent_spot')]:
        row=next(r for r in prices if r[0]==sid)
        result[name]={str(int(period))[:4]+'-'+str(int(period))[4:]:row[i] for i,period in enumerate(calendar)
                      if finite(period) and 190001<=period<=210012 and i<len(row) and finite(row[i])}
    book.close();return result


def shift(period,months):
    year,month=map(int,period.split('-'));total=year*12+month-1+months
    return f'{total//12:04d}-{total%12+1:02d}'


def score(vintages,panel,as_of):
    actual={}
    for key in ('wti_spot','brent_spot'):
        grouped=defaultdict(list)
        for r in panel.get(key,{}).get('observations',[]):
            if r['date'][:7]<as_of.isoformat()[:7]:grouped[r['date'][:7]].append(r['value'])
        actual[key]={p:sum(v)/len(v) for p,v in grouped.items()}
    rows=[]
    for vintage in vintages:
        origin=vintage['published_at'][:7]
        for key,prices in vintage['prices'].items():
            baseline=prices.get(shift(origin,-1))
            if baseline is None:continue
            for horizon in (1,3,6,12):
                target=shift(origin,horizon)
                if target not in prices or target not in actual[key]:continue
                rows.append({'series':key,'origin':origin,'published_at':vintage['published_at'],'target':target,'horizon_months':horizon,
                             'forecast':prices[target],'actual':actual[key][target],'baseline':baseline,'source_url':vintage['source_url'],'sha256':vintage['sha256']})
    summaries=[]
    for key in ('wti_spot','brent_spot'):
        for horizon in (1,3,6,12):
            group=[r for r in rows if r['series']==key and r['horizon_months']==horizon]
            if not group:continue
            summaries.append({'series':key,'horizon_months':horizon,'observations':len(group),
                              'rmse':math.sqrt(sum((r['forecast']-r['actual'])**2 for r in group)/len(group)),
                              'baseline_rmse':math.sqrt(sum((r['baseline']-r['actual'])**2 for r in group)/len(group)),
                              'mae':sum(abs(r['forecast']-r['actual']) for r in group)/len(group)})
    return {'status':'publisher_forecast_evaluation' if rows else 'unavailable','vintages':len(vintages),'rows':rows,'summary':summaries,
            'design':'Fixed 2015–2025 origins; 1/3/6/12 months ahead. Frozen archived forecasts versus completed-month observed spot averages. Baseline is the last completed month in each vintage. Latest revised realized history; no coefficients selected on these outcomes.',
            'limitation':'This evaluates EIA forecasts, not the skill of our stress scenarios or post-2027 price coefficients. Overlapping horizons are not independent observations.'}


def collect(download,panel,directory,as_of):
    from bs4 import BeautifulSoup
    soup=BeautifulSoup(download(ARCHIVE),'html.parser');links=[]
    for row in soup.select('tr'):
        link=next((a['href'] for a in row.select('a[href]') if '_base.xlsx' in a['href']),None)
        stamp=re.search(r'\b(\d{1,2}/\d{1,2}/\d{4})\b',row.get_text(' ',strip=True))
        if not link or not stamp:continue
        published=datetime.strptime(stamp[1],'%m/%d/%Y').date()
        if date(2015,1,1)<=published<=date(2025,12,31) and published<=as_of:links.append((urljoin(ARCHIVE,link),published.isoformat()))
    if not links:raise ValueError('No eligible forecast vintages found')
    def fetch(item):
        url,published=item;cache=directory/(digest(url)+'.json')
        old=load(cache)
        if old:return old
        try:
            raw=download(url);sha=digest(raw)
            directory.mkdir(parents=True,exist_ok=True)
            (directory/(sha+'.xlsx')).write_bytes(raw)
            result={'published_at':published,'source_url':url,'sha256':sha,'prices':parse(raw)}
            save(cache,result);return result
        except Exception as exc:return {'source_url':url,'error':type(exc).__name__}
    with ThreadPoolExecutor(max_workers=4) as pool:results=list(pool.map(fetch,links))
    failures=[r for r in results if r.get('error')]
    result=score([r for r in results if not r.get('error')],panel,as_of)
    result.update(expected_vintages=len(links),failed_vintages=failures,complete=not failures)
    return result
